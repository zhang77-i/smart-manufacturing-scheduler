import math
import random
from bisect import insort
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
HARD_LIMIT = 24 * 3600
DAILY_WINDOW_HOURS = 24
MISSING_OUTSOURCE_HOURS_ASSUMPTION = 24
RANDOM_SEEDS = [29, 37, 62, 63, 86, 103, 119, 125, 147, 167, 179, 181, 221, 229, 258]
SCORE_MODES = ("DUE_WAIT", "DUE_END", "DUE_PROCESS")
OUTPUT_FILENAME = "生产排程日计划.xlsx"


def load_data():
    required_files = [
        "工单基础信息表.xlsx",
        "生产加工时间表.xlsx",
        "外发成本表.xlsx",
        "钻机最早可用时间表.xlsx",
    ]
    missing = [name for name in required_files if not (BASE_DIR / name).exists()]
    if missing:
        raise FileNotFoundError(
            "请把代码与以下输入文件放在同一目录：" + "、".join(missing)
        )

    orders = pd.read_excel(BASE_DIR / "工单基础信息表.xlsx").reset_index(names="订单ID")
    process = pd.read_excel(BASE_DIR / "生产加工时间表.xlsx")
    outsource = pd.read_excel(BASE_DIR / "外发成本表.xlsx")
    machines = pd.read_excel(BASE_DIR / "钻机最早可用时间表.xlsx").reset_index(names="机器索引")

    orders["生产型号"] = orders["生产型号"].astype(str).str.strip()
    process["生产型号"] = process["生产型号"].astype(str).str.strip()
    process["设备组"] = process["设备组"].astype(str).str.strip()
    outsource["生产型号"] = outsource["生产型号"].astype(str).str.strip()
    machines["设备组"] = machines["设备组"].astype(str).str.strip()

    orders["预计到达时间"] = pd.to_datetime(orders["预计到达时间"], errors="raise")
    orders["批次完工时间"] = pd.to_datetime(orders["批次完工时间"], errors="raise")
    machines["钻机最早可用时间"] = pd.to_datetime(machines["钻机最早可用时间"], errors="raise")

    for col in ["产品数量\n（PANEL）", "叠片数"]:
        orders[col] = pd.to_numeric(orders[col], errors="raise").astype(int)

    machines["轴数"] = pd.to_numeric(machines["轴数"], errors="raise").astype(int)
    process["钻孔时间"] = pd.to_numeric(process["钻孔时间"], errors="raise").astype(int)
    outsource["批次外发成本"] = pd.to_numeric(outsource["批次外发成本"], errors="coerce")
    outsource["外发时间"] = pd.to_numeric(outsource["外发时间"], errors="coerce")

    if orders[["生产型号", "产品数量\n（PANEL）", "叠片数", "预计到达时间", "批次完工时间"]].isna().any().any():
        raise ValueError("工单关键字段存在空值")

    if machines[["钻机ID", "设备组", "轴数", "钻机最早可用时间"]].isna().any().any():
        raise ValueError("钻机关键字段存在空值")

    if process[["生产型号", "设备组", "钻孔时间"]].isna().any().any():
        raise ValueError("加工时间关键字段存在空值")

    if (orders["产品数量\n（PANEL）"] <= 0).any() or (orders["叠片数"] <= 0).any():
        raise ValueError("PANEL数量或叠片数存在非正值")

    if (machines["轴数"] <= 0).any() or (process["钻孔时间"] <= 0).any():
        raise ValueError("轴数或钻孔时间存在非正值")

    return orders, process, outsource, machines


def build_process_map(orders, process):
    rows = []

    for row in process.itertuples(index=False):
        for group in str(row.设备组).split("/"):
            rows.append([row.生产型号, group.strip(), int(row.钻孔时间)])

    expanded = pd.DataFrame(rows, columns=["生产型号", "设备组", "钻孔时间"])
    check = expanded.groupby(["生产型号", "设备组"], as_index=False).agg(
        记录数=("钻孔时间", "size"),
        不同时间数=("钻孔时间", "nunique"),
        最小钻孔时间=("钻孔时间", "min"),
        最大钻孔时间=("钻孔时间", "max"),
    )

    conflict = check[check["不同时间数"] > 1].copy()
    relevant = conflict[conflict["生产型号"].isin(set(orders["生产型号"]))].copy()

    if len(relevant) > 0:
        raise ValueError("当前工单涉及同型号同设备组不同钻孔时间的冲突记录")

    clean = expanded.groupby(["生产型号", "设备组"], as_index=False)["钻孔时间"].first()
    process_map = {
        (row.生产型号, row.设备组): int(row.钻孔时间)
        for row in clean.itertuples(index=False)
    }

    return process_map, conflict, relevant


def build_parameters(orders, machines, process_map):
    origin = min(orders["预计到达时间"].min(), machines["钻机最早可用时间"].min())

    release = ((orders["预计到达时间"] - origin).dt.total_seconds()).astype(int).tolist()
    due = ((orders["批次完工时间"] - origin).dt.total_seconds()).astype(int).tolist()
    hard_due = [value + HARD_LIMIT for value in release]
    available = ((machines["钻机最早可用时间"] - origin).dt.total_seconds()).astype(int).tolist()

    models = orders["生产型号"].tolist()
    panels = orders["产品数量\n（PANEL）"].astype(int).tolist()
    stacks = orders["叠片数"].astype(int).tolist()
    groups = machines["设备组"].tolist()
    axes = machines["轴数"].astype(int).tolist()

    candidates = []
    minimum_processing = []

    for i in range(len(orders)):
        alternatives = []

        for m in range(len(machines)):
            key = (models[i], groups[m])

            if key not in process_map:
                continue

            trips = math.ceil(panels[i] / (stacks[i] * axes[m]))
            processing_time = trips * process_map[key]
            alternatives.append((m, int(processing_time), int(trips), int(process_map[key])))

        if not alternatives:
            raise ValueError(f"生产型号 {models[i]} 没有兼容钻机")

        candidates.append(alternatives)
        minimum_processing.append(min(item[1] for item in alternatives))

    return origin, release, due, hard_due, available, candidates, minimum_processing


def structural_check(release, due, hard_due, available, candidates, use_machine_availability):
    unavoidable_outsource = []
    unavoidable_late = []
    rows = []

    for i in range(len(release)):
        earliest_end = min(
            max(release[i], available[m] if use_machine_availability else 0) + processing_time
            for m, processing_time, _, _ in candidates[i]
        )

        if earliest_end > hard_due[i]:
            unavoidable_outsource.append(i)

        if earliest_end > due[i]:
            unavoidable_late.append(i)

        rows.append((i, earliest_end))

    return unavoidable_outsource, unavoidable_late, rows


def first_fit(intervals, release_time, processing_time, available_time, hard_due_time):
    start = max(release_time, available_time)

    for left, right, _ in intervals:
        if start + processing_time <= left:
            return start

        if start < right:
            start = right

        if start + processing_time > hard_due_time:
            return None

    if start + processing_time > hard_due_time:
        return None

    return start


def build_sequence(kind, release, due, minimum_processing, seed=0):
    ids = list(range(len(release)))

    if kind == "EDD":
        ids.sort(key=lambda i: (due[i], release[i], minimum_processing[i], i))
    elif kind == "RELEASE":
        ids.sort(key=lambda i: (release[i], due[i], minimum_processing[i], i))
    elif kind == "SLACK":
        ids.sort(key=lambda i: (due[i] - release[i] - minimum_processing[i], due[i], release[i], i))
    elif kind == "CR":
        ids.sort(key=lambda i: ((due[i] - release[i]) / max(1, minimum_processing[i]), due[i], release[i], i))
    else:
        rng = random.Random(seed)
        alpha, beta, gamma = [rng.uniform(-1, 1) for _ in range(3)]
        ids.sort(
            key=lambda i: (
                due[i]
                + alpha * release[i]
                + beta * minimum_processing[i]
                + gamma * (due[i] - release[i] - minimum_processing[i])
                + rng.random() * 900,
                release[i],
                i,
            )
        )

    return ids


def construct_schedule(
    sequence,
    release,
    due,
    hard_due,
    available,
    candidates,
    machine_count,
    mandatory_outsource,
    use_machine_availability,
    score_mode,
):
    outsource_set = set(mandatory_outsource)
    calendars = [[] for _ in range(machine_count)]
    result = [None] * len(release)

    for i in outsource_set:
        result[i] = {"outsource": 1}

    for i in sequence:
        if i in outsource_set:
            continue

        best = None

        for m, processing_time, trips, drill_time in candidates[i]:
            machine_available = available[m] if use_machine_availability else 0
            start = first_fit(calendars[m], release[i], processing_time, machine_available, hard_due[i])

            if start is None:
                continue

            end = start + processing_time
            tardiness = max(0, end - due[i])
            wait = start - release[i]

            if score_mode == "DUE_WAIT":
                key = (tardiness, wait, end, processing_time)
            elif score_mode == "DUE_END":
                key = (tardiness, end, wait, processing_time)
            else:
                key = (tardiness, processing_time, wait, end)

            if best is None or key < best[0]:
                best = (key, m, processing_time, trips, drill_time, start, end)

        if best is None:
            return None

        _, m, processing_time, trips, drill_time, start, end = best
        insort(calendars[m], (start, end, i))
        result[i] = {
            "outsource": 0,
            "machine": m,
            "processing_time": processing_time,
            "trips": trips,
            "drill_time": drill_time,
            "start": start,
            "end": end,
        }

    return result


def schedule_objective(result, release, due):
    ontime = 0
    tardiness = 0
    waiting = 0

    for i, row in enumerate(result):
        if row["outsource"] == 1:
            continue

        ontime += int(row["end"] <= due[i])
        tardiness += max(0, row["end"] - due[i])
        waiting += row["start"] - release[i]

    return -ontime, tardiness, waiting


def calendars_from_result(result, machine_count):
    calendars = [[] for _ in range(machine_count)]

    for i, row in enumerate(result):
        if row["outsource"] == 0:
            insort(calendars[row["machine"]], (row["start"], row["end"], i))

    return calendars


def two_job_improvement(
    result,
    release,
    due,
    hard_due,
    available,
    candidates,
    machine_count,
    unavoidable_late,
    use_machine_availability,
    max_rounds=10,
):
    unavoidable_late = set(unavoidable_late)

    for _ in range(max_rounds):
        calendars = calendars_from_result(result, machine_count)
        best_objective = schedule_objective(result, release, due)
        best_move = None

        late_jobs = [
            i
            for i, row in enumerate(result)
            if row["outsource"] == 0 and row["end"] > due[i] and i not in unavoidable_late
        ]

        for i in late_jobs:
            old_i = result[i]
            old_machine = old_i["machine"]
            old_interval = (old_i["start"], old_i["end"], i)
            calendars[old_machine].remove(old_interval)

            for m, processing_i, trips_i, drill_i in candidates[i]:
                blocking_jobs = [
                    j
                    for start, end, j in calendars[m]
                    if start < due[i] and end > release[i]
                ]

                for j in blocking_jobs:
                    old_j = result[j]
                    old_j_interval = (old_j["start"], old_j["end"], j)
                    calendars[m].remove(old_j_interval)

                    start_i = first_fit(
                        calendars[m],
                        release[i],
                        processing_i,
                        available[m] if use_machine_availability else 0,
                        hard_due[i],
                    )

                    if start_i is not None and start_i + processing_i <= due[i]:
                        end_i = start_i + processing_i
                        insort(calendars[m], (start_i, end_i, i))
                        best_j = None

                        for machine_j, processing_j, trips_j, drill_j in candidates[j]:
                            start_j = first_fit(
                                calendars[machine_j],
                                release[j],
                                processing_j,
                                available[machine_j] if use_machine_availability else 0,
                                hard_due[j],
                            )

                            if start_j is None:
                                continue

                            end_j = start_j + processing_j
                            key_j = (
                                int(end_j > due[j]),
                                max(0, end_j - due[j]),
                                start_j - release[j],
                                end_j,
                            )

                            if best_j is None or key_j < best_j[0]:
                                best_j = (
                                    key_j,
                                    machine_j,
                                    processing_j,
                                    trips_j,
                                    drill_j,
                                    start_j,
                                    end_j,
                                )

                        if best_j is not None:
                            _, machine_j, processing_j, trips_j, drill_j, start_j, end_j = best_j
                            result[i] = {
                                "outsource": 0,
                                "machine": m,
                                "processing_time": processing_i,
                                "trips": trips_i,
                                "drill_time": drill_i,
                                "start": start_i,
                                "end": end_i,
                            }
                            result[j] = {
                                "outsource": 0,
                                "machine": machine_j,
                                "processing_time": processing_j,
                                "trips": trips_j,
                                "drill_time": drill_j,
                                "start": start_j,
                                "end": end_j,
                            }

                            current_objective = schedule_objective(result, release, due)
                            result[i] = old_i
                            result[j] = old_j

                            if current_objective < best_objective:
                                best_objective = current_objective
                                best_move = (
                                    i,
                                    m,
                                    processing_i,
                                    trips_i,
                                    drill_i,
                                    start_i,
                                    end_i,
                                    j,
                                    machine_j,
                                    processing_j,
                                    trips_j,
                                    drill_j,
                                    start_j,
                                    end_j,
                                )

                        calendars[m].remove((start_i, end_i, i))

                    insort(calendars[m], old_j_interval)

            insort(calendars[old_machine], old_interval)

        if best_move is None:
            break

        (
            i,
            machine_i,
            processing_i,
            trips_i,
            drill_i,
            start_i,
            end_i,
            j,
            machine_j,
            processing_j,
            trips_j,
            drill_j,
            start_j,
            end_j,
        ) = best_move

        result[i] = {
            "outsource": 0,
            "machine": machine_i,
            "processing_time": processing_i,
            "trips": trips_i,
            "drill_time": drill_i,
            "start": start_i,
            "end": end_i,
        }
        result[j] = {
            "outsource": 0,
            "machine": machine_j,
            "processing_time": processing_j,
            "trips": trips_j,
            "drill_time": drill_j,
            "start": start_j,
            "end": end_j,
        }

    return result


def solve_heuristic(
    release,
    due,
    hard_due,
    available,
    candidates,
    minimum_processing,
    machine_count,
    mandatory_outsource,
    unavoidable_late,
    use_machine_availability,
):
    configurations = []

    for kind in ["EDD", "RELEASE", "SLACK", "CR"]:
        for score_mode in SCORE_MODES:
            configurations.append((kind, score_mode, 0))

    for seed in RANDOM_SEEDS:
        for score_mode in SCORE_MODES:
            configurations.append(("RANDOM", score_mode, seed))

    feasible_solutions = []

    for kind, score_mode, seed in configurations:
        sequence = build_sequence(kind, release, due, minimum_processing, seed)
        result = construct_schedule(
            sequence,
            release,
            due,
            hard_due,
            available,
            candidates,
            machine_count,
            mandatory_outsource,
            use_machine_availability,
            score_mode,
        )

        if result is None:
            continue

        feasible_solutions.append(
            (schedule_objective(result, release, due), kind, score_mode, seed, result)
        )

    if not feasible_solutions:
        raise RuntimeError("启发式未找到可行排程")

    feasible_solutions.sort(key=lambda item: item[0])
    improve_pool = feasible_solutions[:32]

    best_result = None
    best_objective = None
    best_method = None

    for _, kind, score_mode, seed, result in improve_pool:
        candidate_result = [row.copy() for row in result]
        candidate_result = two_job_improvement(
            candidate_result,
            release,
            due,
            hard_due,
            available,
            candidates,
            machine_count,
            unavoidable_late,
            use_machine_availability,
        )
        current_objective = schedule_objective(candidate_result, release, due)

        if best_result is None or current_objective < best_objective:
            best_result = candidate_result
            best_objective = current_objective
            best_method = f"{kind}-{score_mode}-SEED{seed}"

    return best_result, best_objective, best_method


def solve_with_outsource_escalation(
    release,
    due,
    hard_due,
    available,
    candidates,
    minimum_processing,
    machine_count,
    mandatory_outsource,
    unavoidable_late,
    use_machine_availability,
):
    outsource_set = set(mandatory_outsource)
    escalation_order = sorted(
        (i for i in range(len(release)) if i not in outsource_set),
        key=lambda i: (
            hard_due[i] - release[i] - minimum_processing[i],
            due[i] - release[i] - minimum_processing[i],
            -minimum_processing[i],
            i,
        ),
    )

    while True:
        try:
            result, objective, method = solve_heuristic(
                release,
                due,
                hard_due,
                available,
                candidates,
                minimum_processing,
                machine_count,
                sorted(outsource_set),
                unavoidable_late,
                use_machine_availability,
            )
            method = f"OUTSOURCE-{len(outsource_set)}|{method}"
            return result, objective, method, sorted(outsource_set)
        except RuntimeError:
            if not escalation_order:
                raise RuntimeError("即使全部订单允许外发，启发式仍未找到可行方案")
            outsource_set.add(escalation_order.pop(0))


def build_outsource_map(outsource):
    valid = outsource.dropna(subset=["批次外发成本", "外发时间"]).copy()
    grouped = valid.groupby("生产型号", as_index=False).agg(
        最低外发成本=("批次外发成本", "min"),
        外发时间=("外发时间", "min"),
        成本记录数=("批次外发成本", "size"),
        不同成本数=("批次外发成本", "nunique"),
    )

    mapping = {
        row.生产型号: (
            float(row.最低外发成本),
            int(row.外发时间),
            int(row.成本记录数),
            int(row.不同成本数),
        )
        for row in grouped.itertuples(index=False)
    }

    return mapping, grouped


def build_result(
    scenario,
    result,
    orders,
    machines,
    origin,
    release,
    due,
    hard_due,
    outsource_map,
):
    rows = []

    for i, schedule in enumerate(result):
        order = orders.loc[i]

        if schedule["outsource"] == 1:
            cost_info = outsource_map.get(order["生产型号"])
            outsource_hours = (
                cost_info[1]
                if cost_info is not None
                else MISSING_OUTSOURCE_HOURS_ASSUMPTION
            )
            cost = cost_info[0] if cost_info is not None else math.nan
            cost_records = cost_info[2] if cost_info is not None else 0
            different_costs = cost_info[3] if cost_info is not None else 0
            outsource_time_source = "外发成本表" if cost_info is not None else "缺失型号统一假设"
            outsource_time_verified = int(cost_info is not None)
            start = release[i]
            end = release[i] + outsource_hours * 3600
            machine_index = math.nan
            machine_id = ""
            group = "外发"
            axes = 0
            trips = 0
            drill_time = 0
            processing_time = 0
        else:
            machine_index = schedule["machine"]
            machine_id = machines.loc[machine_index, "钻机ID"]
            group = machines.loc[machine_index, "设备组"]
            axes = int(machines.loc[machine_index, "轴数"])
            trips = int(schedule["trips"])
            drill_time = int(schedule["drill_time"])
            processing_time = int(schedule["processing_time"])
            start = int(schedule["start"])
            end = int(schedule["end"])
            cost = 0.0
            cost_records = 0
            different_costs = 0
            outsource_hours = 0
            outsource_time_source = "不适用"
            outsource_time_verified = 1

        rows.append(
            {
                "方案": scenario,
                "订单ID": int(i),
                "生产型号": order["生产型号"],
                "PANEL数量": int(order["产品数量\n（PANEL）"]),
                "叠片数": int(order["叠片数"]),
                "预计到达时间": order["预计到达时间"],
                "批次完工时间": order["批次完工时间"],
                "24小时截止时间": origin + pd.Timedelta(seconds=hard_due[i]),
                "机器索引": machine_index,
                "钻机ID": machine_id,
                "设备组": group,
                "轴数": axes,
                "加工趟数": trips,
                "单趟钻孔时间秒": drill_time,
                "总加工时间秒": processing_time,
                "计划开始时间": origin + pd.Timedelta(seconds=start),
                "计划结束时间": origin + pd.Timedelta(seconds=end),
                "等待时间小时": 0.0 if schedule["outsource"] == 1 else (start - release[i]) / 3600,
                "流动时间小时": (end - release[i]) / 3600,
                "逾期时间小时": max(0, end - due[i]) / 3600,
                "是否按批次交期完成": int(end <= due[i]),
                "是否24小时内完成": int(end <= hard_due[i]),
                "是否外发": int(schedule["outsource"]),
                "外发时间小时": outsource_hours,
                "外发时间来源": outsource_time_source,
                "外发时间是否有数据": outsource_time_verified,
                "最低外发成本": cost,
                "外发成本记录数": cost_records,
                "不同外发成本数": different_costs,
                "外发成本是否缺失": int(schedule["outsource"] == 1 and cost_info is None),
            }
        )

    return pd.DataFrame(rows)


def validate_result(
    result,
    orders,
    machines,
    process_map,
    use_machine_availability,
):
    errors = []

    if len(result) != len(orders):
        errors.append("结果工单数与输入工单数不一致")

    if result["订单ID"].nunique() != len(orders):
        errors.append("订单ID存在重复或缺失")

    for row in result.itertuples(index=False):
        i = int(row.订单ID)

        if int(row.是否外发) == 1:
            if int(row.是否24小时内完成) != 1:
                errors.append(f"订单{i}:外发后超过24小时")
            continue

        machine_index = int(row.机器索引)
        order = orders.loc[i]
        machine = machines.loc[machine_index]
        key = (order["生产型号"], machine["设备组"])

        if key not in process_map:
            errors.append(f"订单{i}:机器设备组不兼容")
            continue

        trips = math.ceil(
            int(order["产品数量\n（PANEL）"])
            / (int(order["叠片数"]) * int(machine["轴数"]))
        )
        processing_time = trips * int(process_map[key])

        if int(row.加工趟数) != trips:
            errors.append(f"订单{i}:加工趟数错误")

        if int(row.总加工时间秒) != processing_time:
            errors.append(f"订单{i}:加工时间错误")

        if row.计划开始时间 < order["预计到达时间"]:
            errors.append(f"订单{i}:早于订单到达时间")

        if use_machine_availability and row.计划开始时间 < machine["钻机最早可用时间"]:
            errors.append(f"订单{i}:早于钻机最早可用时间")

        if row.计划结束时间 - row.计划开始时间 != pd.Timedelta(seconds=processing_time):
            errors.append(f"订单{i}:开始结束时间与加工时间不一致")

        if row.计划结束时间 > row.预计到达时间 + pd.Timedelta(hours=24):
            errors.append(f"订单{i}:厂内加工超过24小时")

    inside = result[result["是否外发"] == 0].copy()

    for machine_id, jobs in inside.groupby("钻机ID"):
        jobs = jobs.sort_values("计划开始时间").reset_index(drop=True)

        for k in range(1, len(jobs)):
            if jobs.loc[k, "计划开始时间"] < jobs.loc[k - 1, "计划结束时间"]:
                errors.append(f"钻机{machine_id}:加工区间重叠")

    return errors


def build_metrics(
    scenario,
    result,
    orders,
    machines,
    origin,
    use_machine_availability,
    unavoidable_outsource,
    unavoidable_late,
    validation_errors,
    method,
):
    inside = result[result["是否外发"] == 0].copy()
    outsourced = result[result["是否外发"] == 1].copy()
    horizon_start = orders["预计到达时间"].min()
    horizon_end = orders["预计到达时间"].max() + pd.Timedelta(hours=24)
    horizon_seconds = int((horizon_end - origin).total_seconds())

    if use_machine_availability:
        capacity = sum(
            max(
                0,
                horizon_seconds
                - int((row["钻机最早可用时间"] - origin).total_seconds()),
            )
            for _, row in machines.iterrows()
        )
    else:
        capacity = len(machines) * horizon_seconds

    rolling_utilization = inside["总加工时间秒"].sum() / capacity

    daily_end = horizon_start + pd.Timedelta(hours=DAILY_WINDOW_HOURS)
    daily_busy_seconds = 0
    for row in inside.itertuples(index=False):
        overlap_start = max(row.计划开始时间, horizon_start)
        overlap_end = min(row.计划结束时间, daily_end)
        daily_busy_seconds += max(0, int((overlap_end - overlap_start).total_seconds()))

    if use_machine_availability:
        daily_capacity = sum(
            max(
                0,
                int(
                    (
                        daily_end
                        - max(horizon_start, row["钻机最早可用时间"])
                    ).total_seconds()
                ),
            )
            for _, row in machines.iterrows()
        )
    else:
        daily_capacity = len(machines) * DAILY_WINDOW_HOURS * 3600

    daily_utilization = daily_busy_seconds / daily_capacity
    known_cost = outsourced["最低外发成本"].dropna().sum()

    return {
        "方案": scenario,
        "求解方法": method,
        "求解状态": "HEURISTIC_FEASIBLE" if len(validation_errors) == 0 else "INVALID_SCHEDULE",
        "订单总数": len(result),
        "厂内订单数": len(inside),
        "外发订单数": int(result["是否外发"].sum()),
        "结构性最少外发下界": len(unavoidable_outsource),
        "按批次交期完成数": int(result["是否按批次交期完成"].sum()),
        "订单满足率": float(result["是否按批次交期完成"].mean()),
        "结构性必迟订单数": len(unavoidable_late),
        "订单满足率理论上界": 1 - len(unavoidable_late) / len(result),
        "距离理论上界订单数": len(result) - len(unavoidable_late) - int(result["是否按批次交期完成"].sum()),
        "24小时完成率（含缺失外发时长假设）": float(result["是否24小时内完成"].mean()),
        "厂内平均等待时间小时": float(inside["等待时间小时"].mean()),
        "厂内平均流动时间小时": float(inside["流动时间小时"].mean()),
        "厂内平均逾期时间小时": float(inside["逾期时间小时"].mean()),
        "滚动窗口开始": origin,
        "滚动窗口结束": horizon_end,
        "滚动窗口小时": horizon_seconds / 3600,
        "滚动完工窗口设备稼动率": float(rolling_utilization),
        "24小时日窗口设备稼动率": float(daily_utilization),
        "按最低报价口径已知外发成本": float(known_cost),
        "外发成本缺失订单数": int(outsourced["外发成本是否缺失"].sum()),
        "外发时长缺失订单数": int((outsourced["外发时间是否有数据"] == 0).sum()),
        "排程验证错误数": len(validation_errors),
    }


def build_daily_plan(result):
    plan = result[result["是否外发"] == 0].copy()
    plan = plan.sort_values(["钻机ID", "计划开始时间", "订单ID"]).reset_index(drop=True)
    plan["计划日期"] = plan["计划开始时间"].dt.normalize()
    plan["生产顺序"] = plan.groupby("钻机ID").cumcount() + 1
    plan["批次交期"] = plan["批次完工时间"]
    plan["是否预计按期"] = plan["是否按批次交期完成"].map({1: "是", 0: "否"})

    return plan[
        [
            "计划日期",
            "钻机ID",
            "生产顺序",
            "订单ID",
            "生产型号",
            "PANEL数量",
            "叠片数",
            "加工趟数",
            "计划开始时间",
            "计划结束时间",
            "批次交期",
            "是否预计按期",
        ]
    ]


def build_outsource_plan(q3_result, q4_result):
    q3_ids = set(q3_result.loc[q3_result["是否外发"] == 1, "订单ID"].astype(int))
    q4_ids = set(q4_result.loc[q4_result["是否外发"] == 1, "订单ID"].astype(int))

    if q3_ids == q4_ids:
        outsourced = q4_result[q4_result["是否外发"] == 1].copy()
        outsourced["适用方案"] = "Q3、Q4"
    else:
        frames = []
        for label, result in [("Q3", q3_result), ("Q4", q4_result)]:
            part = result[result["是否外发"] == 1].copy()
            part["适用方案"] = label
            frames.append(part)
        outsourced = pd.concat(frames, ignore_index=True)

    outsourced["预计完成时间"] = outsourced["计划结束时间"]

    def make_note(row):
        time_missing = int(row["外发时间是否有数据"]) == 0
        cost_missing = int(row["外发成本是否缺失"]) == 1
        if time_missing and cost_missing:
            return f"外发时长及成本数据缺失，暂按{MISSING_OUTSOURCE_HOURS_ASSUMPTION}小时假设"
        if time_missing:
            return f"外发时长数据缺失，暂按{MISSING_OUTSOURCE_HOURS_ASSUMPTION}小时假设"
        if cost_missing:
            return "外发成本数据缺失"
        return "外发数据完整"

    outsourced["备注"] = outsourced.apply(make_note, axis=1)
    outsourced = outsourced.sort_values(["适用方案", "订单ID"]).reset_index(drop=True)

    return outsourced[
        [
            "适用方案",
            "订单ID",
            "生产型号",
            "PANEL数量",
            "预计到达时间",
            "24小时截止时间",
            "外发时间小时",
            "预计完成时间",
            "最低外发成本",
            "备注",
        ]
    ].rename(columns={"最低外发成本": "预计外发成本"})


def build_indicator_comparison(q2_summary, q3_summary, q4_summary):
    return pd.DataFrame(
        [
            {
                "方案": "Q2 基础模型",
                "条件变化": "全部订单必须厂内生产",
                "运行结论": "结构性不可行",
                "订单总数": q2_summary["订单总数"],
                "厂内订单数": math.nan,
                "外发订单数": 0,
                "按批次交期完成数": math.nan,
                "订单满足率": math.nan,
                "24小时完成率": math.nan,
                "24小时日窗口设备稼动率": math.nan,
                "厂内平均等待时间（小时）": math.nan,
                "校验错误数": math.nan,
            },
            {
                "方案": "Q3 允许外发",
                "条件变化": "无法在24小时内完成的订单允许外发",
                "运行结论": "启发式可行",
                "订单总数": q3_summary["订单总数"],
                "厂内订单数": q3_summary["厂内订单数"],
                "外发订单数": q3_summary["外发订单数"],
                "按批次交期完成数": q3_summary["按批次交期完成数"],
                "订单满足率": q3_summary["订单满足率"],
                "24小时完成率": q3_summary["24小时完成率（含缺失外发时长假设）"],
                "24小时日窗口设备稼动率": q3_summary["24小时日窗口设备稼动率"],
                "厂内平均等待时间（小时）": q3_summary["厂内平均等待时间小时"],
                "校验错误数": q3_summary["排程验证错误数"],
            },
            {
                "方案": "Q4 最终方案",
                "条件变化": "在Q3基础上加入钻机最早可用时间",
                "运行结论": "启发式可行",
                "订单总数": q4_summary["订单总数"],
                "厂内订单数": q4_summary["厂内订单数"],
                "外发订单数": q4_summary["外发订单数"],
                "按批次交期完成数": q4_summary["按批次交期完成数"],
                "订单满足率": q4_summary["订单满足率"],
                "24小时完成率": q4_summary["24小时完成率（含缺失外发时长假设）"],
                "24小时日窗口设备稼动率": q4_summary["24小时日窗口设备稼动率"],
                "厂内平均等待时间（小时）": q4_summary["厂内平均等待时间小时"],
                "校验错误数": q4_summary["排程验证错误数"],
            },
        ]
    )


def build_method_notes():
    algorithm = pd.DataFrame(
        [
            ["1. 可行机器筛选", "按生产型号、设备组筛选兼容钻机", "压缩订单—机器候选集合"],
            ["2. 加工时间计算", "ceil(PANEL数量/(叠片数×轴数))×单趟钻孔时间", "得到异构机器加工时间"],
            ["3. 结构性检验", "检查订单独占最快兼容机器时能否在24小时内完成", "证明Q2不可行并得到最少外发下界"],
            ["4. 多起点构造", "EDD、SLACK、CR、到达时间及随机序列", "生成多组候选排程"],
            ["5. First-Fit分配", "遍历机器时间轴，并组合DUE_WAIT、DUE_END、DUE_PROCESS三种评分策略", "确定机器、开始时间和顺序"],
            ["6. 局部改进", "对可改进迟交订单尝试换机与两订单重排", "提高按期完成数并降低等待"],
            ["7. 独立校验", "复查兼容性、加工时长、到达/可用时间、24小时与区间重叠", "校验失败则终止导出"],
        ],
        columns=["步骤", "代码实现", "建模作用"],
    )

    assumptions = pd.DataFrame(
        [
            ["加工时间", "业务已确认钻孔时间为单趟标准时间；加工趟数=ceil(PANEL数量/(叠片数×轴数))，实际加工时间=加工趟数×钻孔时间"],
            ["订单批次", "订单不可拆分且不可抢占；同一订单的全部PANEL只分配到一台兼容钻机连续加工"],
            ["24小时硬约束", "每个订单的截止时间=预计到达时间+24小时"],
            ["订单满足率", "计划完成时间不晚于批次完工时间的订单占全部订单比例"],
            ["设备稼动率", "主表采用从最早到达时间起24小时的固定日窗口；Q4扣除机器初始不可用时间"],
            ["外发资源", "假设外发无并发容量限制，订单到达后可立即外发；如供应商存在产能限制，应增加外发容量约束"],
            ["缺失外发时长", f"缺失型号暂按{MISSING_OUTSOURCE_HOURS_ASSUMPTION}小时，相关完成率为条件性结论"],
            ["求解状态", "启发式可行不等于全局最优；仅对结构性不可行与上下界给出证明"],
        ],
        columns=["口径或假设", "说明"],
    )
    return algorithm, assumptions


NAVY = "1F4E78"
LIGHT_BLUE = "EAF2F8"
YELLOW = "FFF2CC"
WHITE = "FFFFFF"
TEXT = "1F2937"


def add_title(ws, title, subtitle, max_col):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    ws.cell(1, 1, title)
    ws.cell(1, 1).fill = PatternFill("solid", fgColor=NAVY)
    ws.cell(1, 1).font = Font(name="微软雅黑", size=16, bold=True, color=WHITE)
    ws.cell(1, 1).alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_col)
    ws.cell(2, 1, subtitle)
    ws.cell(2, 1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    ws.cell(2, 1).font = Font(name="微软雅黑", size=10, color=TEXT)
    ws.cell(2, 1).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 30


def format_table_sheet(
    ws,
    title,
    subtitle,
    widths,
    freeze_panes="A5",
):
    header_row = 4
    add_title(ws, title, subtitle, len(widths))
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = freeze_panes
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(widths))}{ws.max_row}"

    for cell in ws[header_row]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="微软雅黑", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 34

    for col_index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_index)].width = width

    date_only = {"计划日期"}
    datetime_headers = {
        "计划开始时间",
        "计划结束时间",
        "批次交期",
        "预计到达时间",
        "24小时截止时间",
        "预计完成时间",
        "单独占用最快兼容机器的最早完成时间",
    }
    percent_headers = {"订单满足率", "24小时完成率", "24小时日窗口设备稼动率"}

    for row in range(header_row + 1, ws.max_row + 1):
        ws.row_dimensions[row].height = 20
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row, col)
            cell.font = Font(name="微软雅黑", size=9, color=TEXT)
            cell.alignment = Alignment(vertical="center")
            header = ws.cell(header_row, col).value
            if header in date_only:
                cell.number_format = "yyyy-mm-dd"
            elif header in datetime_headers:
                cell.number_format = "yyyy-mm-dd hh:mm"
            elif header in percent_headers:
                cell.number_format = "0.00%"
            elif header in {"厂内平均等待时间（小时）", "最快完成所需小时", "超过24小时（小时）"}:
                cell.number_format = "0.000"
            elif header == "预计外发成本":
                cell.number_format = "#,##0.00"
            if header in {"备注", "条件变化", "运行结论"}:
                cell.alignment = Alignment(vertical="center", wrap_text=True)

    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def format_method_sheet(ws, algorithm_rows, assumption_rows):
    ws.sheet_view.showGridLines = False
    add_title(
        ws,
        "建模与启发式算法说明",
        "展示从问题抽象、结构性检验、多起点构造到局部改进和结果校验的完整实现主线。",
        3,
    )
    ws.freeze_panes = "A6"
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 56
    ws.column_dimensions["C"].width = 42

    ws.merge_cells("A4:C4")
    ws["A4"] = "启发式算法步骤"
    ws["A4"].fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    ws["A4"].font = Font(name="微软雅黑", bold=True, color=NAVY)
    ws["A4"].alignment = Alignment(vertical="center")

    algorithm_header = 5
    for cell in ws[algorithm_header]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="微软雅黑", bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in range(algorithm_header + 1, algorithm_header + 1 + algorithm_rows):
        ws.row_dimensions[row].height = 34
        for cell in ws[row]:
            cell.font = Font(name="微软雅黑", size=9, color=TEXT)
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    assumption_section = algorithm_header + algorithm_rows + 3
    ws.merge_cells(start_row=assumption_section, start_column=1, end_row=assumption_section, end_column=3)
    ws.cell(assumption_section, 1, "关键口径与假设")
    ws.cell(assumption_section, 1).fill = PatternFill("solid", fgColor=YELLOW)
    ws.cell(assumption_section, 1).font = Font(name="微软雅黑", bold=True, color="7F6000")

    assumption_header = assumption_section + 1
    for cell in ws[assumption_header][:2]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="微软雅黑", bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in range(assumption_header + 1, assumption_header + 1 + assumption_rows):
        ws.row_dimensions[row].height = 34
        for col in range(1, 3):
            cell = ws.cell(row, col)
            cell.font = Font(name="微软雅黑", size=9, color=TEXT)
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def export_excel(output, q4_plan, outsource_plan, comparison, q3_plan, q2_infeasible, algorithm, assumptions):
    with pd.ExcelWriter(output, engine="openpyxl", datetime_format="yyyy-mm-dd hh:mm") as writer:
        q4_plan.to_excel(writer, sheet_name="Q4最终日生产计划", index=False, startrow=3)
        outsource_plan.to_excel(writer, sheet_name="外发计划", index=False, startrow=3)
        comparison.to_excel(writer, sheet_name="方案指标对比", index=False, startrow=3)
        q3_plan.to_excel(writer, sheet_name="Q3日生产计划", index=False, startrow=3)
        q2_infeasible.to_excel(writer, sheet_name="Q2不可行订单", index=False, startrow=3)
        algorithm.to_excel(writer, sheet_name="算法说明与假设", index=False, startrow=4)
        assumption_section = 5 + len(algorithm) + 3
        assumptions.to_excel(
            writer,
            sheet_name="算法说明与假设",
            index=False,
            startrow=assumption_section,
        )

        workbook = writer.book

        format_table_sheet(
            workbook["Q4最终日生产计划"],
            "企业最终日生产计划（Q4）",
            "已考虑订单外发与钻机最早可用时间；按钻机ID、计划开始时间排序，生产顺序可直接执行。",
            [13, 10, 10, 10, 17, 12, 10, 10, 20, 20, 20, 14],
            freeze_panes="D5",
        )
        format_table_sheet(
            workbook["外发计划"],
            "订单外发计划",
            "列明无法在24小时内由厂内完成的订单；缺失外发数据已在备注中明确标出。",
            [12, 10, 17, 12, 20, 20, 14, 20, 16, 42],
        )
        format_table_sheet(
            workbook["方案指标对比"],
            "Q2—Q4生产排程方案对比",
            "Q2用于证明基础条件不可行；Q3引入外发；Q4进一步加入机器最早可用时间并形成最终计划。",
            [18, 34, 20, 12, 12, 12, 16, 14, 14, 22, 22, 14],
        )

        format_table_sheet(
            workbook["Q3日生产计划"],
            "Q3日生产计划（允许外发）",
            "在基础排程模型上增加外发选择；按钻机ID、计划开始时间排序。",
            [13, 10, 10, 10, 17, 12, 10, 10, 20, 20, 20, 14],
            freeze_panes="D5",
        )
        format_table_sheet(
            workbook["Q2不可行订单"],
            "Q2基础模型：结构性不可行证明",
            "以下订单即使单独占用最快兼容钻机，仍无法在到达后24小时内完成，因此全部厂内生产无可行解。",
            [10, 17, 12, 10, 20, 20, 24, 18, 18],
        )
        format_method_sheet(workbook["算法说明与假设"], len(algorithm), len(assumptions))


def main():
    orders, process, outsource, machines = load_data()
    process_map, _, _ = build_process_map(orders, process)
    origin, release, due, hard_due, available, candidates, minimum_processing = build_parameters(
        orders, machines, process_map
    )
    outsource_map, _ = build_outsource_map(outsource)

    q2_out, q2_late, q2_earliest = structural_check(
        release, due, hard_due, available, candidates, False
    )
    q4_out, q4_late, _ = structural_check(
        release, due, hard_due, available, candidates, True
    )

    q3_schedule, q3_objective, q3_method, _ = solve_with_outsource_escalation(
        release,
        due,
        hard_due,
        available,
        candidates,
        minimum_processing,
        len(machines),
        q2_out,
        q2_late,
        False,
    )

    q4_schedule, q4_objective, q4_method, _ = solve_with_outsource_escalation(
        release,
        due,
        hard_due,
        available,
        candidates,
        minimum_processing,
        len(machines),
        q4_out,
        q4_late,
        True,
    )

    q3_result = build_result(
        "Q3_允许外发",
        q3_schedule,
        orders,
        machines,
        origin,
        release,
        due,
        hard_due,
        outsource_map,
    )
    q4_result = build_result(
        "Q4_外发+机器最早可用时间",
        q4_schedule,
        orders,
        machines,
        origin,
        release,
        due,
        hard_due,
        outsource_map,
    )

    q3_errors = validate_result(q3_result, orders, machines, process_map, False)
    q4_errors = validate_result(q4_result, orders, machines, process_map, True)

    if q3_errors or q4_errors:
        raise RuntimeError(
            "排程校验失败："
            + "；".join([*q3_errors, *q4_errors])
        )

    q2_rows = []

    for i in q2_out:
        earliest_end = q2_earliest[i][1]
        q2_rows.append(
            {
                "订单ID": i,
                "生产型号": orders.loc[i, "生产型号"],
                "PANEL数量": int(orders.loc[i, "产品数量\n（PANEL）"]),
                "叠片数": int(orders.loc[i, "叠片数"]),
                "预计到达时间": orders.loc[i, "预计到达时间"],
                "24小时截止时间": orders.loc[i, "预计到达时间"] + pd.Timedelta(hours=24),
                "单独占用最快兼容机器的最早完成时间": origin + pd.Timedelta(seconds=earliest_end),
                "最快完成所需小时": (earliest_end - release[i]) / 3600,
                "超过24小时（小时）": (earliest_end - release[i]) / 3600 - 24,
            }
        )

    q2_infeasible = pd.DataFrame(q2_rows)

    q2_summary = {"订单总数": len(orders)}

    q3_summary = build_metrics(
        "Q3_允许外发",
        q3_result,
        orders,
        machines,
        origin,
        False,
        q2_out,
        q2_late,
        q3_errors,
        q3_method,
    )
    q4_summary = build_metrics(
        "Q4_外发+机器最早可用时间",
        q4_result,
        orders,
        machines,
        origin,
        True,
        q4_out,
        q4_late,
        q4_errors,
        q4_method,
    )

    q3_plan = build_daily_plan(q3_result)
    q4_plan = build_daily_plan(q4_result)
    outsource_plan = build_outsource_plan(q3_result, q4_result)
    comparison = build_indicator_comparison(q2_summary, q3_summary, q4_summary)
    algorithm, assumptions = build_method_notes()
    output = BASE_DIR / OUTPUT_FILENAME

    export_excel(
        output,
        q4_plan,
        outsource_plan,
        comparison,
        q3_plan,
        q2_infeasible,
        algorithm,
        assumptions,
    )

    print("\n=== 生产排程运行结果 ===")
    print(f"Q2：结构性不可行；至少 {len(q2_out)} 单必须外发")
    print(
        f"Q3：外发 {q3_summary['外发订单数']} 单，"
        f"订单满足率 {q3_summary['订单满足率']:.2%}，"
        f"24小时日窗口稼动率 {q3_summary['24小时日窗口设备稼动率']:.2%}，"
        f"平均等待 {q3_summary['厂内平均等待时间小时']:.3f} 小时"
    )
    print(
        f"Q4：外发 {q4_summary['外发订单数']} 单，"
        f"订单满足率 {q4_summary['订单满足率']:.2%}，"
        f"24小时日窗口稼动率 {q4_summary['24小时日窗口设备稼动率']:.2%}，"
        f"平均等待 {q4_summary['厂内平均等待时间小时']:.3f} 小时"
    )
    print(f"Q3目标值: {q3_objective}")
    print(f"Q4目标值: {q4_objective}")
    print(f"结果文件: {output}")


if __name__ == "__main__":
    main()
